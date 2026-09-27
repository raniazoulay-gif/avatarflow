"""Word extraction.

DOCX: paragraphs, tables, headers and footers via python-docx.
DOC (legacy binary): best effort via `antiword`, then LibreOffice conversion.
LIMITATION: if neither tool is installed, .doc files cannot be read reliably
and are marked as errors (NOT processed silently).
"""

from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile


class WordError(Exception):
    pass


def extract_docx_text(data: bytes) -> str:
    try:
        import docx

        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise WordError(f"Broken/unreadable DOCX: {type(exc).__name__}") from exc

    parts: list[str] = []
    for section in d.sections:
        for hf in (section.header, section.first_page_header, section.footer,
                   section.first_page_footer):
            try:
                parts.extend(p.text for p in hf.paragraphs if p.text.strip())
                for t in hf.tables:
                    parts.extend(_table_rows(t))
            except Exception:
                continue
    parts.extend(p.text for p in d.paragraphs if p.text.strip())
    for t in d.tables:
        parts.extend(_table_rows(t))
    return "\n".join(parts)


def _table_rows(table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


def doc_support() -> str:
    if shutil.which("antiword"):
        return "antiword"
    if shutil.which("soffice") or shutil.which("libreoffice"):
        return "libreoffice"
    return "none"


def extract_doc_text(data: bytes, timeout: int = 60) -> str:
    tool = doc_support()
    if tool == "none":
        raise WordError("LIMITATION: legacy .doc not supported (install antiword or LibreOffice)")
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "in.doc")
        with open(path, "wb") as f:
            f.write(data)
        if tool == "antiword":
            r = subprocess.run(["antiword", "-w", "0", path], capture_output=True, timeout=timeout)
            if r.returncode == 0 and r.stdout.strip():
                return r.stdout.decode("utf-8", errors="replace")
            if not (shutil.which("soffice") or shutil.which("libreoffice")):
                raise WordError("Could not read .doc with antiword")
        binary = shutil.which("soffice") or shutil.which("libreoffice")
        if binary is None:
            raise WordError("Could not read .doc (LibreOffice not installed)")
        r = subprocess.run(
            [binary, "--headless", "--convert-to", "txt:Text", "--outdir", tmp, path],
            capture_output=True, timeout=timeout,
        )
        out = os.path.join(tmp, "in.txt")
        if r.returncode != 0 or not os.path.exists(out):
            raise WordError("Could not convert .doc with LibreOffice")
        with open(out, encoding="utf-8", errors="replace") as f:
            return f.read()
