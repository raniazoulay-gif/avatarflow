"""Daily Excel workbook: REPORT, SUMMARY, SUPPLIERS, QUALITY sheets."""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import DETAIL_COLUMNS, ReportData

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
YES_FILL = PatternFill("solid", fgColor="E2EFDA")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")
ERROR_FILL = PatternFill("solid", fgColor="F8CBAD")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def _header(ws, cols: list[str], row: int = 1) -> None:
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autowidth(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_width,
                                                                           max(10, length + 2))


def write_excel(data: ReportData, reports_dir: str) -> str:
    os.makedirs(reports_dir, exist_ok=True)
    path = os.path.join(reports_dir, f"invoice_report_{data.report_date.isoformat()}.xlsx")
    wb = Workbook()

    # SUMMARY first so it opens by default
    ws = wb.active
    ws.title = "SUMMARY"
    r = 1
    if data.dry_run:
        ws.cell(row=r, column=1, value="⚠ DRY RUN MODE — No emails were forwarded.").font = \
            Font(bold=True, size=14, color="9C5700")
        ws.cell(row=r, column=1).fill = WARN_FILL
        r += 2
    ws.cell(row=r, column=1, value=f"Daily Invoice Report – {data.report_date.isoformat()}"
            ).font = Font(bold=True, size=12)
    r += 1
    _header(ws, ["Metric", "Value"], row=r)
    for k, v in data.summary.items():
        r += 1
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    _autowidth(ws)

    # REPORT (detail rows)
    wd = wb.create_sheet("REPORT")
    _header(wd, DETAIL_COLUMNS)
    for i, row in enumerate(data.rows, 2):
        for j, col in enumerate(DETAIL_COLUMNS, 1):
            wd.cell(row=i, column=j, value=row.get(col))
        status = row.get("Processing Status", "")
        fill = None
        if status == "ERROR":
            fill = ERROR_FILL
        elif row.get("Would Forward") == "YES":
            fill = YES_FILL
        elif "REVIEW" in status or row.get("Possible Invoice") == "YES":
            fill = REVIEW_FILL
        if fill:
            for j in range(1, len(DETAIL_COLUMNS) + 1):
                wd.cell(row=i, column=j).fill = fill
    for col in ("L", "M", "N", "O"):
        for c in wd[col][1:]:
            c.number_format = "0.00"
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = wd.dimensions
    _autowidth(wd)

    # SUPPLIERS
    wsu = wb.create_sheet("SUPPLIERS")
    cols = ["Supplier", "Email", "Invoices", "High Confidence", "Review", "Would Forward", "Errors"]
    _header(wsu, cols)
    for i, sup in enumerate(data.suppliers, 2):
        for j, c in enumerate(cols, 1):
            wsu.cell(row=i, column=j, value=sup.get(c))
    _autowidth(wsu)

    # QUALITY (dry-run evaluation metrics)
    wq = wb.create_sheet("QUALITY")
    _header(wq, ["Metric", "Value"])
    for i, (metric, value) in enumerate(data.quality.items(), 2):
        wq.cell(row=i, column=1, value=metric)
        wq.cell(row=i, column=2, value=value)
    _autowidth(wq)

    wb.save(path)
    return path
