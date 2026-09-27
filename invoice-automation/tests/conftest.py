from __future__ import annotations

import base64
import io
import json
import os
import re
import shutil
from typing import Any

import pytest

from src.classification.ai_classifier import AIClassifier
from src.config.settings import Settings
from src.database.repository import Database
from src.documents.extractor import DocumentExtractor
from src.gmail.forwarder import Forwarder
from src.gmail.labels import LabelManager
from src.processor import Processor
from src.utils import retry

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

INVOICE_LINES = [
    "Example Supplies Ltd",
    "VAT Reg No: 514000000",
    "TAX INVOICE",
    "Invoice No: INV-2026-0917",
    "Invoice Date: 15/09/2026",
    "Due Date: 15/10/2026",
    "Bill To: Customer Co.",
    "Consulting services              1,000.00",
    "Subtotal: 1,000.00 ILS",
    "VAT 18%: 180.00 ILS",
    "Total Amount Due: 1,180.00 ILS",
]
NON_INVOICE_LINES = [
    "Autumn Newsletter",
    "Dear friends, we are excited to share our latest product news.",
    "Join our webinar next week to learn about new features and community events.",
    "Thank you for being part of our community. Unsubscribe any time.",
]


# ---------------------------------------------------------------- documents
def make_text_pdf(lines: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in lines:
        c.drawString(60, y, line)
        y -= 22
    c.save()
    return buf.getvalue()


def make_scanned_pdf(lines: list[str]) -> bytes:
    """Image-only PDF (no text layer) - simulates a scanned invoice."""
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (1700, 1300), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, 40) if os.path.exists(FONT) else ImageFont.load_default()
    y = 60
    for line in lines:
        draw.text((80, y), line, fill="black", font=font)
        y += 70
    buf = io.BytesIO()
    img.save(buf, format="PDF", resolution=150)
    return buf.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None,
              header: str | None = None) -> bytes:
    import docx

    d = docx.Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, val in enumerate(row):
                t.cell(i, j).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- fake gmail
class ForwardAPICalled(AssertionError):
    pass


class FakeGmail:
    """In-memory Gmail API double."""

    def __init__(self) -> None:
        self.messages: dict[str, dict[str, Any]] = {}
        self.attachments: dict[tuple[str, str], bytes] = {}
        self.raw: dict[str, bytes] = {}
        self.labels: dict[str, str] = {}
        self.applied: dict[str, list[str]] = {}
        self.sent: list[bytes] = []
        self.fail_get: set[str] = set()
        self.fail_send = False
        self.forbid_send = False  # raise if send is ever called

    def add_message(self, mid: str, *, sender: str = "Supplier <billing@supplier.co.il>",
                    subject: str = "Invoice", body: str = "Please find attached.",
                    attachments: list[tuple[str, str, bytes]] | None = None,
                    internal_ms: int = 1790000000000) -> None:
        parts = [{"partId": "0", "mimeType": "text/plain", "filename": "",
                  "body": {"data": base64.urlsafe_b64encode(body.encode()).decode(),
                           "size": len(body)}}]
        from email.message import EmailMessage

        em = EmailMessage()
        em["From"] = sender
        em["To"] = "me@example.com"
        em["Subject"] = subject
        em.set_content(body)
        for i, (fname, mime, data) in enumerate(attachments or [], 1):
            aid = f"att-{mid}-{i}"
            self.attachments[(mid, aid)] = data
            parts.append({"partId": str(i), "mimeType": mime, "filename": fname,
                          "body": {"attachmentId": aid, "size": len(data)}})
            mt, _, st = mime.partition("/")
            em.add_attachment(data, maintype=mt, subtype=st, filename=fname)
        self.raw[mid] = em.as_bytes()
        self.messages[mid] = {
            "id": mid, "threadId": f"t-{mid}", "internalDate": str(internal_ms),
            "payload": {"mimeType": "multipart/mixed",
                        "headers": [{"name": "From", "value": sender},
                                    {"name": "Subject", "value": subject}],
                        "parts": parts},
        }

    # --- GmailAPI protocol
    def get_profile(self):
        return {"emailAddress": "me@example.com"}

    def list_message_ids(self, query: str, max_results: int = 500):
        return list(reversed(list(self.messages)))

    def get_message(self, mid):
        if mid in self.fail_get:
            raise ConnectionError("Gmail API unavailable")
        return self.messages[mid]

    def get_raw_message(self, mid):
        return self.raw[mid]

    def get_attachment(self, mid, aid):
        return self.attachments[(mid, aid)]

    def list_labels(self):
        return [{"name": n, "id": i} for n, i in self.labels.items()]

    def create_label(self, name):
        self.labels[name] = f"Label_{len(self.labels) + 1}"
        return {"id": self.labels[name], "name": name}

    def add_labels(self, mid, label_ids):
        names = {v: k for k, v in self.labels.items()}
        self.applied.setdefault(mid, []).extend(names[i] for i in label_ids)

    def send_raw(self, raw, thread_id=None):
        if self.forbid_send:
            raise ForwardAPICalled("Gmail send API was called!")
        if self.fail_send:
            raise ConnectionError("send failed")
        self.sent.append(raw)
        return {"id": f"sent-{len(self.sent)}"}

    def watch(self, topic):
        return {"historyId": "1", "expiration": "1790000000000"}


# ---------------------------------------------------------------- fake AI
class FakeBackend:
    """Keyword-driven stand-in for the LLM."""

    def __init__(self, fail: bool = False, bad_json_times: int = 0) -> None:
        self.fail = fail
        self.bad_json_times = bad_json_times
        self.calls = 0
        self.prompts: list[str] = []

    def complete(self, system: str, user: str) -> str:
        self.calls += 1
        self.prompts.append(user)
        if self.fail:
            raise ConnectionError("AI provider down")
        if self.bad_json_times > 0:
            self.bad_json_times -= 1
            return "Sure! Here is the answer: {is_invoice: yes"
        doc = user.split("--- DOCUMENT TEXT (truncated) ---", 1)[-1].lower()
        if re.search(r"invoice|חשבונית", doc):
            return json.dumps({
                "is_invoice": True, "confidence": 0.97, "invoice_type": "tax_invoice",
                "supplier_name": "Example Supplies Ltd", "invoice_number": "INV-2026-0917",
                "invoice_date": "2026-09-15", "due_date": "2026-10-15", "currency": "ILS",
                "subtotal": 1000, "vat": 180, "total": 1180,
                "reason": "Contains invoice number, VAT and total.",
            })
        return json.dumps({"is_invoice": False, "confidence": 0.96, "invoice_type": "none",
                           "supplier_name": None, "invoice_number": None, "invoice_date": None,
                           "due_date": None, "currency": None, "subtotal": None, "vat": None,
                           "total": None, "reason": "Marketing newsletter."})


# ---------------------------------------------------------------- wiring
@pytest.fixture(autouse=True)
def _no_sleep(monkeypatch):
    monkeypatch.setattr(retry, "sleep", lambda s: None)


def make_settings(**overrides) -> Settings:
    base = dict(
        _env_file=None,
        source_gmail_account="me@example.com",
        target_gmail_account="books@example.com",
        database_url="sqlite:///:memory:",
        dry_run=True, auto_forward_enabled=False, production_confirmation=False,
        retry_base_delay_seconds=0, retry_max_attempts=3,
        ai_api_key="",
    )
    base.update(overrides)
    return Settings(**base)


PRODUCTION = dict(dry_run=False, auto_forward_enabled=True, production_confirmation=True)


class Harness:
    def __init__(self, settings: Settings, backend: FakeBackend | None = None,
                 ocr_pdf_fn=None, reports_dir: str | None = None,
                 notify: bool = False) -> None:
        self.settings = settings
        self.gmail = FakeGmail()
        self.db = Database(settings.database_url)
        self.db.create_all()
        self.backend = backend if backend is not None else FakeBackend()
        self.ai = AIClassifier(self.backend, max_attempts=3, base_delay=0)
        kwargs = {"retry_attempts": 3, "retry_base_delay": 0}
        if ocr_pdf_fn is not None:
            kwargs["ocr_pdf_fn"] = ocr_pdf_fn
        self.extractor = DocumentExtractor(**kwargs)
        self.labels = LabelManager(self.gmail)
        self.forwarder = Forwarder(settings, self.gmail, self.labels)
        from src.reports.notifications import Notifier

        # Notifications are off in the harness unless a test opts in, so the
        # "no send in DRY RUN" assertions keep meaning "no forward".
        self.notifier = Notifier(settings, self.gmail) if notify else None
        self.processor = Processor(settings, self.db, self.gmail, self.extractor, self.ai,
                                   self.forwarder, self.labels, self.notifier)

    def email(self, mid: str):
        with self.db.repo() as repo:
            e = repo.get_email(mid)
            if e is not None:
                _ = e.classification, e.forward, e.attachments
            return e


@pytest.fixture
def harness_factory():
    return Harness


def tesseract_ok() -> bool:
    return shutil.which("tesseract") is not None and shutil.which("pdftoppm") is not None
