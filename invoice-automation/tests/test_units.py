"""Unit tests: rules, decision engine, AI JSON parsing, retry, logging, reports, watcher."""

from __future__ import annotations

import logging
from datetime import datetime, timedelta

import openpyxl
import pytest

from src.classification import rule_engine
from src.classification.ai_classifier import InvalidAIResponse, parse_response
from src.classification.decision_engine import decide, final_confidence
from src.database.models import EmailStatus
from src.documents.extractor import detect_file_type, is_relevant
from src.gmail.reader import parse_message
from src.gmail.watcher import Watcher
from src.reports.email_report import build_body, generate_daily_report
from src.system_check import run_system_check
from src.utils import retry
from src.utils.logging_setup import RedactingFilter, redact

from .conftest import INVOICE_LINES, NON_INVOICE_LINES, PRODUCTION, Harness, make_settings, make_text_pdf


def test_rule_engine_scores():
    hi = rule_engine.evaluate("\n".join(INVOICE_LINES))
    lo = rule_engine.evaluate("\n".join(NON_INVOICE_LINES))
    assert hi.score >= 0.85
    assert hi.invoice_number == "INV-2026-0917"
    assert lo.score <= 0.25


def test_rule_engine_hebrew_variants():
    text = "חשבונית מס מס' 1234\nתאריך: 01/09/2026\nסה״כ לתשלום: 590.00 ש\"ח\nמע״מ 90.00\nח.פ. 512345678"
    r = rule_engine.evaluate(text)
    assert all(r.indicators.values()), r.indicators
    assert r.currency == "ILS"


def test_final_confidence_weights_configurable():
    from src.classification.ai_classifier import AIClassification

    rules = rule_engine.evaluate("\n".join(INVOICE_LINES))
    ai = AIClassification(is_invoice=True, confidence=1.0, reason="")
    s1 = make_settings(ai_weight=0.7, rule_weight=0.3)
    s2 = make_settings(ai_weight=0.0, rule_weight=1.0)
    assert final_confidence(ai, rules, s1) == pytest.approx(0.7 + 0.3 * rules.score)
    assert final_confidence(ai, rules, s2) == pytest.approx(rules.score)
    assert final_confidence(None, rules, s1) <= 0.89


def test_decision_bands():
    s = make_settings()
    assert decide(0.95, "a@b.com", s).status == EmailStatus.DRY_RUN_WOULD_FORWARD
    assert decide(0.80, "a@b.com", s).status == EmailStatus.REVIEW
    assert decide(0.30, "a@b.com", s).status == EmailStatus.NOT_INVOICE
    p = make_settings(**PRODUCTION)
    assert decide(0.95, "a@b.com", p).status == EmailStatus.FORWARDED
    assert decide(0.90, "a@b.com", p).status == EmailStatus.FORWARDED
    assert decide(0.8999, "a@b.com", p).status == EmailStatus.REVIEW


def test_parse_ai_json_valid_and_fenced():
    r = parse_response('```json\n{"is_invoice": true, "confidence": 0.9, "total": "1,180.00"}\n```')
    assert r.is_invoice and r.total == 1180.0


@pytest.mark.parametrize("bad", ["no json here", '{"is_invoice": true}',
                                 '{"is_invoice": true, "confidence": 7}', "{broken"])
def test_parse_ai_json_invalid(bad):
    with pytest.raises(InvalidAIResponse):
        parse_response(bad)


def test_retry_exponential_backoff(monkeypatch):
    delays = []
    monkeypatch.setattr(retry, "sleep", delays.append)
    calls = {"n": 0}

    def flaky():
        calls["n"] += 1
        if calls["n"] < 5:
            raise ConnectionError()
        return "ok"

    assert retry.retry_call(flaky, max_attempts=5, base_delay=1) == "ok"
    assert delays == [1, 2, 4, 8]


def test_retry_gives_up():
    with pytest.raises(ConnectionError):
        retry.retry_call(lambda: (_ for _ in ()).throw(ConnectionError()), max_attempts=3,
                         base_delay=0)


def test_log_redaction():
    assert "sk-ant-abc" not in redact("key sk-ant-abc123XYZ")
    assert "secret123" not in redact("refresh_token=secret123")
    rec = logging.LogRecord("x", logging.INFO, "", 0, "token %s", ("mysecretvalue",), None)
    RedactingFilter(["mysecretvalue"]).filter(rec)
    assert "mysecretvalue" not in rec.getMessage()


def test_file_type_detection():
    assert detect_file_type("a.pdf", None) == "pdf"
    assert detect_file_type("x.bin", None, b"%PDF-1.7") == "pdf"
    assert detect_file_type("a.DOCX", None) == "docx"
    assert not is_relevant("signature.png", "image/png")
    assert not is_relevant("tracking.gif", "image/gif")
    assert is_relevant("invoice_scan.jpg", "image/jpeg")
    assert not is_relevant("data.csv", "text/csv")


def test_reader_parses_gmail_payload():
    h = Harness(make_settings())
    h.gmail.add_message("m1", sender="Acme Ltd <Billing@Acme.com>", subject="Inv",
                        attachments=[("a.pdf", "application/pdf", b"%PDF")])
    p = parse_message(h.gmail.messages["m1"])
    assert p.sender_email == "billing@acme.com"
    assert p.sender_name == "Acme Ltd"
    assert p.attachments[0].filename == "a.pdf"
    assert p.body_text.startswith("Please find")


def test_daily_report_dry_run(tmp_path):
    s = make_settings(reports_dir=str(tmp_path))
    h = Harness(s)
    h.gmail.add_message("m1", attachments=[("i.pdf", "application/pdf",
                                            make_text_pdf(INVOICE_LINES))])
    h.gmail.add_message("m2", sender="News <n@shop.com>", subject="Newsletter",
                        attachments=[("n.pdf", "application/pdf",
                                      make_text_pdf(NON_INVOICE_LINES))])
    h.gmail.add_message("m3", subject="Invoice for August", attachments=[])
    h.processor.process_many(["m1", "m2", "m3"])
    now = datetime.now(s.tz)
    # The report for day D covers (D-1 18:00, D 18:00]; pick the one containing "now".
    d = now.date() if (now.hour, now.minute) < (s.report_hour, s.report_minute) \
        else now.date() + timedelta(days=1)
    res = generate_daily_report(s, h.db, h.gmail, d, send=True)
    assert res["sent"] is True
    summary = res["summary"]
    assert summary["Total Emails"] == 3
    assert summary["Would Be Forwarded"] == 1
    assert summary["Actually Forwarded"] == 0
    assert summary["Possible False Negatives"] == 1
    wb = openpyxl.load_workbook(res["excel"])
    assert set(wb.sheetnames) == {"SUMMARY", "REPORT", "SUPPLIERS", "QUALITY"}
    assert "DRY RUN" in wb["SUMMARY"]["A1"].value
    headers = [c.value for c in wb["REPORT"][1]]
    assert "Would Forward" in headers and "Actually Forwarded" in headers
    rows = {r[5]: r for r in wb["REPORT"].iter_rows(min_row=2, values_only=True)}
    assert rows["m1"][headers.index("Actually Forwarded")] == "NO – DRY RUN"
    assert rows["m1"][headers.index("Would Forward")] == "YES"
    mail = h.gmail.sent[0].decode("utf-8", errors="replace")
    assert "Daily Invoice Automation Report" in mail
    assert "To: me@example.com" in mail
    assert "books@example.com" not in mail


def test_report_body_dry_run_banner():
    from src.reports.data import ReportData

    d = ReportData(report_date=datetime(2026, 9, 27).date(), dry_run=True, summary={
        "Total Emails": 1, "Invoices Detected": 1, "Would Be Forwarded": 1,
        "Actually Forwarded": 0, "Review Required": 0, "Possible False Negatives": 0,
        "Errors": 0}, quality={})
    body = build_body(d)
    assert body.startswith("⚠ DRY RUN MODE\nNo emails were forwarded.")
    assert "DRY RUN MODE: YES" in body


def test_dry_run_period_never_switches_to_production(tmp_path):
    from datetime import date, timedelta

    from src.reports.email_report import dry_run_progress

    s = make_settings(reports_dir=str(tmp_path), dry_run_days=7)
    h = Harness(s)
    start = date(2026, 9, 1)
    assert "day 1 of 7" in dry_run_progress(h.db, s, start)
    msg = dry_run_progress(h.db, s, start + timedelta(days=10))
    assert "REMAINS in DRY RUN" in msg
    assert s.dry_run is True and s.mode == "DRY_RUN"


def test_watcher_poll_only_new_and_idempotent():
    s = make_settings()
    h = Harness(s)
    w = Watcher(s, h.db, h.gmail, h.processor)
    h.gmail.add_message("m1", attachments=[("i.pdf", "application/pdf",
                                            make_text_pdf(INVOICE_LINES))])
    assert w.poll_once() == {"m1": EmailStatus.DRY_RUN_WOULD_FORWARD}
    assert w.poll_once() == {}
    assert w.start_watch().startswith("NOT CONFIGURED")


def test_system_check_not_configured_is_honest(monkeypatch):
    s = make_settings(gmail_client_id="", ai_api_key="")
    items = {i.name: i for i in run_system_check(s, verify_remote=True)}
    assert items["Gmail"].status == "NOT CONFIGURED"
    assert items["AI"].status == "NOT CONFIGURED"
    assert items["Mode"].status == "DRY RUN"
    assert items["Target Gmail"].status == "CONFIGURED"
    assert "NOT VERIFIED" in items["Target Gmail"].detail


def test_health_endpoint():
    from fastapi.testclient import TestClient

    from src.api.server import create_app
    from src.app_context import AppContext

    s = make_settings()
    ctx = AppContext(s, connect_gmail=False)
    client = TestClient(create_app(ctx))
    body = client.get("/health").json()
    assert body["mode"] == "DRY_RUN"
    assert body["gmail_status"] == "NOT CONFIGURED"
    assert body["ai_status"] == "NOT CONFIGURED"
    assert body["database_status"] == "OK"
    for k in ("last_processed_email", "last_report", "last_error", "monitoring"):
        assert k in body
    assert client.post("/gmail/push?token=x").status_code == 403


def test_scheduler_uses_explicit_timezone():
    from src.app_context import AppContext
    from src.scheduler.jobs import build_scheduler

    s = make_settings(timezone="Asia/Jerusalem", daily_report_time="18:00")
    sched = build_scheduler(AppContext(s, connect_gmail=False))
    job = sched.get_job("daily_report")
    assert str(job.trigger.timezone) == "Asia/Jerusalem"
    assert "hour='18'" in str(job.trigger) and "minute='0'" in str(job.trigger)


def test_public_pages_for_google_consent_screen():
    from fastapi.testclient import TestClient

    from src.api.server import create_app
    from src.app_context import AppContext

    s = make_settings(public_contact_email="owner@example.com")
    client = TestClient(create_app(AppContext(s, connect_gmail=False)))
    for path, marker in (("/", "Invoice Automation"), ("/privacy", "Limited Use"),
                         ("/terms", "Terms of Service")):
        r = client.get(path)
        assert r.status_code == 200 and "text/html" in r.headers["content-type"]
        assert marker in r.text
        assert "owner@example.com" in r.text
    # Public pages never leak system state.
    assert "total_processed" not in client.get("/").text


def test_gmail_auth_error_is_explained_without_secrets():
    from google.auth.exceptions import RefreshError

    from src.gmail.auth import describe_auth_error

    exc = RefreshError("invalid_grant: Token has been expired or revoked.",
                       {"error": "invalid_grant",
                        "error_description": "Token has been expired or revoked."})
    msg = describe_auth_error(exc)
    assert "invalid_grant" in msg and "OAuth Playground" in msg
    assert describe_auth_error(RefreshError("invalid_client: Unauthorized")).startswith(
        "RefreshError: invalid_client")
    assert describe_auth_error(ConnectionError("x")) == "ConnectionError"


def test_gmail_credential_paste_mistakes_detected():
    from src.gmail.auth import credential_format_problems

    s = make_settings(gmail_client_id="<apps.googleusercontent.com>",
                      gmail_client_secret="GOCSPX-abc", gmail_refresh_token="<1//04abc>")
    problems = " ".join(credential_format_problems(s))
    assert "GMAIL_CLIENT_ID contains" in problems
    assert "GMAIL_REFRESH_TOKEN should start with 1//" in problems
    assert "04abc" not in problems  # values are never echoed
    ok = make_settings(gmail_client_id=" 1-a.apps.googleusercontent.com\n",
                       gmail_client_secret="GOCSPX-abc", gmail_refresh_token="1//04abc ")
    assert credential_format_problems(ok) == []  # whitespace is stripped by settings


def test_system_check_reports_gmail_reason():
    from google.auth.exceptions import RefreshError

    class Broken:
        def get_profile(self):
            raise RefreshError("invalid_grant: Token has been expired or revoked.")

    s = make_settings(gmail_client_id="1-a.apps.googleusercontent.com",
                      gmail_client_secret="GOCSPX-x", gmail_refresh_token="1//x")
    items = {i.name: i for i in run_system_check(s, gmail_client=Broken())}
    assert items["Gmail"].status == "ERROR"
    assert "invalid_grant" in items["Gmail"].detail


def test_paste_mistake_reason_reaches_system_check():
    s = make_settings(gmail_client_id="1-a.apps.googleusercontent.com",
                      gmail_client_secret="GOCSPX-x", gmail_refresh_token="<1//secretvalue>")
    items = {i.name: i for i in run_system_check(s)}  # real build_credentials path
    assert items["Gmail"].status == "ERROR"
    assert "GMAIL_REFRESH_TOKEN" in items["Gmail"].detail
    assert "secretvalue" not in items["Gmail"].detail
